"""CV and cover-letter models built only from verified evidence, plus HTML, DOCX and PDF output."""

from __future__ import annotations

import base64
import hashlib
import io
import json
import os
import re
import shutil
import subprocess
import tempfile
from html import escape
from pathlib import Path
from typing import Dict, List, Optional

from . import evidence

ROOT = Path(__file__).resolve().parent.parent
PHOTO = ROOT / "data" / "photo.png"

TEMPLATES = {
    "professional-operations": {"accent": "#1f3a5f", "font": "'Source Serif Pro', Georgia, serif", "label": "Professional operations"},
    "modern-executive": {"accent": "#2d2d2d", "font": "'Inter', 'Helvetica Neue', Arial, sans-serif", "label": "Modern executive"},
    "research-policy": {"accent": "#3b4a2f", "font": "'Iowan Old Style', Palatino, Georgia, serif", "label": "Research and policy"},
    "environment-projects": {"accent": "#24584a", "font": "'Inter', 'Helvetica Neue', Arial, sans-serif", "label": "Environment and projects"},
    "international-mobility": {"accent": "#7a3b1d", "font": "'Source Serif Pro', Georgia, serif", "label": "International mobility"},
}
TEMPLATE_FOR = {
    "mobility": "international-mobility", "executive_admin": "modern-executive", "operations": "professional-operations", "documentation": "professional-operations",
    "people_ops": "professional-operations", "projects": "environment-projects", "environment": "environment-projects", "research": "research-policy",
    "education_admin": "modern-executive",
}
PRIORITY_TAGS = {
    "mobility": ["immigration", "documentation", "deadlines", "communication", "international"],
    "executive_admin": ["organisation", "deadlines", "documentation", "stakeholders", "crm", "confidentiality"],
    "operations": ["deadlines", "organisation", "documentation", "stakeholders", "crm"],
    "documentation": ["documentation", "compliance", "confidentiality", "deadlines"],
    "people_ops": ["hr_systems", "documentation", "communication", "confidentiality"],
    "projects": ["documentation", "deadlines", "stakeholders", "organisation"],
    "environment": ["documentation", "deadlines", "stakeholders"],
    "research": ["documentation", "deadlines", "organisation"],
    "education_admin": ["communication", "international", "documentation", "organisation"],
}
EDUCATION_FIRST = {"environment", "research", "projects"}


def _photo_data_uri() -> str:
    if PHOTO.exists():
        return "data:image/png;base64," + base64.b64encode(PHOTO.read_bytes()).decode()
    return ""


def _choose_bullets(role_id: str, wanted: List[str], limit: int) -> List[Dict]:
    items = [e for e in evidence.usable() if e.get("parent") == role_id]
    def rank(item):
        tags = item.get("tags", [])
        score = 0
        for position, tag in enumerate(wanted):
            if tag in tags:
                score += 10 - position
        return -score
    items.sort(key=rank)
    return [{"text": item["fact"], "evidence": [item["id"]]} for item in items[:limit]]


def cv_model(job: Dict, analysis: Dict, *, pages: int = 1, template: Optional[str] = None) -> Dict:
    data = evidence.profile()
    family = analysis.get("variant") or "operations"
    matrix_tags = [row["concept"] for row in analysis.get("matrix", []) if row["match"] in ("strong", "transferable")]
    wanted = list(dict.fromkeys(matrix_tags + PRIORITY_TAGS.get(family, [])))
    identity = data["identity"]
    roles = []
    for role in data["experience"]:
        limit = (5 if pages == 2 else 4) if role["id"] == "exp_1" else (4 if pages == 2 else 3)
        roles.append({
            "id": role["id"], "title": role["title"], "employer": role["employer"], "note": role.get("employer_note", ""),
            "place": role["place"], "dates": role["dates"], "bullets": _choose_bullets(role["id"], wanted, limit),
        })
    education = []
    for item in data["education"]:
        if item["id"] == "edu_school" and family not in ("education_admin", "executive_admin"):
            continue
        lines = item.get("lines", []) if (family in EDUCATION_FIRST or pages == 2) else item.get("lines", [])[:1]
        evidence_ids = [e["id"] for e in evidence.usable() if e.get("parent") == item["id"]]
        education.append({"id": item["id"], "credential": item["credential"], "school": item["school"], "place": item.get("place", ""), "dates": item["dates"], "lines": lines, "evidence": evidence_ids})
    skills = []
    for skill in data["skills"]:
        ids = [i for i in skill["evidence"] if (evidence.by_id(i) or {}).get("confirmed")]
        if ids:
            skills.append({"label": skill["label"], "evidence": ids})
    if family not in ("environment", "research", "projects"):
        skills = [s for s in skills if "GIS" not in s["label"] and "Statistics" not in s["label"]]
    languages = [{"label": l["cv"], "evidence": [l["id"]]} for l in data["languages"]]
    summary = data["summaries"].get(family) or data["summaries"]["operations"]
    sections = [
        {"type": "summary", "title": "Profile", "text": summary, "evidence": ["r1_ind", "msc_env", "ba_eu", "lang_en", "lang_nl"]},
        {"type": "experience", "title": "Experience", "entries": roles},
        {"type": "education", "title": "Education", "entries": education},
        {"type": "skills", "title": "Skills", "items": skills},
        {"type": "languages", "title": "Languages", "items": languages},
    ]
    if family in EDUCATION_FIRST:
        sections[1], sections[2] = sections[2], sections[1]
    return {
        "version": 1,
        "template": template or TEMPLATE_FOR.get(family, "professional-operations"),
        "pages": pages,
        "photo": bool(data["preferences"].get("photo", True)),
        "header": {
            "name": identity["name"], "location": f"{identity['city']}, {identity['country']}", "phone": identity["phone"], "email": identity["email"],
            "line": "Dutch and British national · EU right to work",
        },
        "sections": sections,
        "job": {"title": job.get("title"), "company": job.get("company")},
    }


def cover_model(job: Dict, analysis: Dict) -> Dict:
    data = evidence.profile()
    family = analysis.get("variant") or "operations"
    title = job.get("title") or "the role"
    company = job.get("company") or "your team"
    strong = [row for row in analysis.get("matrix", []) if row["match"] == "strong" and row["concept"] not in ("english", "degree_any", "office_tools")]
    focus = ", ".join(row["label"].lower() for row in strong[:2])
    opening = f"I am applying for the {title} role at {company}."
    if focus:
        opening += f" The posting puts weight on {focus}, which is the kind of work I have been doing."
    opening += " " + (data["cover"].get(family) or data["cover"]["operations"])
    picks: List[Dict] = []
    wanted = [row["concept"] for row in strong] + PRIORITY_TAGS.get(family, [])
    for tag in wanted:
        for item in evidence.usable():
            if tag in item.get("tags", []) and (item["parent"].startswith("exp_") or item["parent"] == "edu_msc") and item not in picks:
                picks.append(item)
                break
        if len(picks) >= 3:
            break
    sentences = []
    academic = {
        "msc_env": "My MSc covered environmental and resource management, including biodiversity and ecosystem services.",
        "msc_methods": "My MSc gave me academic training in qualitative and quantitative research, statistics, and data processing and visualisation.",
        "msc_gis": "My MSc included academic GIS work in QGIS.",
        "skill_reports": "My studies involved regular desk research and report writing.",
    }
    for item in picks:
        if item["id"] in academic:
            sentences.append(academic[item["id"]])
            continue
        role = next((r for r in data["experience"] if r["id"] == item["parent"]), None)
        if not role:
            continue
        where = f"At {role['employer']}, as {role['title']}," if re.search(r"sales", role["title"], re.I) else f"At {role['employer']}"
        fact = item["fact"].rstrip(".")
        sentences.append(f"{where} I {fact[0].lower() + fact[1:]}.")
    evidence_paragraph = " ".join(sentences) or "My last full role was case-based operations work: documents, deadlines and accurate records."
    geo = analysis.get("geo", {})
    where = "I live in Thessaloniki" if geo.get("greece_remote") == "onsite" else "I am based in Thessaloniki and work comfortably remotely"
    closing = (
        f"{where}, and as a Dutch and British national I can work in Greece and the EU without sponsorship. English and Dutch are my native languages; "
        "my Greek is basic (A2) and improving. I would welcome a conversation about how I could help, and I can share more detail about the casework I have handled."
    )
    paragraphs = [
        {"text": opening, "evidence": ["r1_ind"]},
        {"text": evidence_paragraph, "evidence": [p["id"] for p in picks]},
        {"text": closing, "evidence": ["nationality", "lang_en", "lang_nl", "lang_gr"]},
    ]
    return {
        "version": 1,
        "greeting": "Dear hiring team,",
        "paragraphs": paragraphs,
        "signoff": "Kind regards,",
        "name": data["identity"]["name"],
        "contact": f"{data['identity']['phone']} · {data['identity']['email']}",
        "ai_policy": analysis.get("ai_policy", "unknown"),
    }


def model_text(model: Dict) -> str:
    parts = []
    for section in model.get("sections", []):
        parts.append(section.get("text", ""))
        for entry in section.get("entries", []):
            parts.append(f"{entry.get('title', '')} {entry.get('employer', '')} {entry.get('credential', '')}")
            parts.extend(b["text"] for b in entry.get("bullets", []))
            parts.extend(entry.get("lines", []))
        parts.extend(i["label"] for i in section.get("items", []))
    for paragraph in model.get("paragraphs", []):
        parts.append(paragraph["text"])
    return "\n".join(p for p in parts if p)


def lint_model(model: Dict) -> List[Dict]:
    return evidence.lint(model_text(model), kind="cover" if "paragraphs" in model else "cv")


def cv_html(model: Dict, *, for_print: bool = True, photo_uri: Optional[str] = None) -> str:
    style = TEMPLATES.get(model.get("template"), TEMPLATES["professional-operations"])
    header = model["header"]
    photo = photo_uri if photo_uri is not None else _photo_data_uri()
    photo_html = f'<img class="photo" src="{photo}" alt="">' if model.get("photo") and photo else ""
    body = []
    for section in model["sections"]:
        body.append(f'<section class="block {section["type"]}"><h2>{escape(section["title"])}</h2>')
        if section["type"] == "summary":
            body.append(f"<p>{escape(section['text'])}</p>")
        elif section["type"] in ("experience",):
            for entry in section["entries"]:
                note = f' <span class="note">({escape(entry["note"])})</span>' if entry.get("note") else ""
                bullets = "".join(f"<li>{escape(b['text'])}</li>" for b in entry["bullets"])
                body.append(f'<div class="entry"><div class="row"><h3>{escape(entry["title"])}</h3><span class="dates">{escape(entry["dates"])}</span></div>'
                            f'<p class="meta">{escape(entry["employer"])}{note} · {escape(entry["place"])}</p><ul>{bullets}</ul></div>')
        elif section["type"] == "education":
            for entry in section["entries"]:
                lines = "".join(f"<li>{escape(line)}</li>" for line in entry.get("lines", []))
                body.append(f'<div class="entry"><div class="row"><h3>{escape(entry["credential"])}</h3><span class="dates">{escape(entry["dates"])}</span></div>'
                            f'<p class="meta">{escape(entry["school"])}{(" · " + escape(entry["place"])) if entry.get("place") else ""}</p>{("<ul>" + lines + "</ul>") if lines else ""}</div>')
        else:
            items = " · ".join(escape(i["label"]) for i in section["items"])
            body.append(f"<p>{items}</p>")
        body.append("</section>")
    page_rule = "@page { size: A4; margin: 14mm 15mm; }"
    return f"""<!DOCTYPE html><html lang="en"><head><meta charset="utf-8"><title>{escape(header['name'])} CV</title>
<style>
{page_rule}
* {{ box-sizing: border-box; }}
body {{ margin: 0; font-family: {style['font']}; color: #1b1b1b; font-size: {'10.2pt' if model.get('pages') == 1 else '10.6pt'}; line-height: 1.38; }}
.cv {{ max-width: 180mm; margin: 0 auto; padding: {'0' if for_print else '12mm 0'}; }}
header {{ display: flex; gap: 14px; align-items: center; border-bottom: 2px solid {style['accent']}; padding-bottom: 8px; margin-bottom: 10px; }}
.photo {{ width: 27mm; height: 33mm; object-fit: cover; border-radius: 3px; }}
h1 {{ margin: 0; font-size: 21pt; color: {style['accent']}; font-weight: 600; letter-spacing: .01em; }}
.contact {{ margin: 3px 0 0; color: #333; }}
h2 {{ font-size: 10.5pt; text-transform: uppercase; letter-spacing: .08em; color: {style['accent']}; margin: 11px 0 5px; border-bottom: 1px solid #ddd; padding-bottom: 2px; }}
h3 {{ font-size: 10.8pt; margin: 0; }}
.row {{ display: flex; justify-content: space-between; gap: 8px; align-items: baseline; }}
.dates {{ color: #555; white-space: nowrap; font-size: 9.6pt; }}
.meta {{ margin: 1px 0 3px; color: #444; }}
.note {{ color: #666; }}
ul {{ margin: 2px 0 6px; padding-left: 16px; }}
li {{ margin: 1px 0; }}
.entry {{ break-inside: avoid; }}
p {{ margin: 2px 0 6px; }}
</style></head><body><div class="cv">
<header>{photo_html}<div><h1>{escape(header['name'])}</h1>
<p class="contact">{escape(header['location'])} · {escape(header['phone'])} · {escape(header['email'])}<br>{escape(header['line'])}</p></div></header>
{''.join(body)}
</div></body></html>"""


def cover_html(model: Dict, job: Dict) -> str:
    paragraphs = "".join(f"<p>{escape(p['text'])}</p>" for p in model["paragraphs"])
    return f"""<!DOCTYPE html><html lang="en"><head><meta charset="utf-8"><title>Cover letter</title>
<style>@page {{ size: A4; margin: 22mm 22mm; }} body {{ font-family: Georgia, serif; font-size: 11pt; line-height: 1.5; color: #1b1b1b; max-width: 165mm; margin: 0 auto; }}</style></head><body>
<p>{escape(model['name'])}<br>{escape(model['contact'])}</p>
<p>Re: {escape(job.get('title') or '')}, {escape(job.get('company') or '')}</p>
<p>{escape(model['greeting'])}</p>{paragraphs}<p>{escape(model['signoff'])}<br>{escape(model['name'])}</p></body></html>"""


def docx_bytes(model: Dict) -> bytes:
    from docx import Document
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Mm, Pt, RGBColor

    style = TEMPLATES.get(model.get("template"), TEMPLATES["professional-operations"])
    accent = RGBColor.from_string(style["accent"].lstrip("#").upper())
    document = Document()
    section = document.sections[0]
    section.page_height, section.page_width = Mm(297), Mm(210)
    section.top_margin = section.bottom_margin = Mm(14)
    section.left_margin = section.right_margin = Mm(15)
    normal = document.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    header = model["header"]
    table = document.add_table(rows=1, cols=2)
    table.autofit = False
    left, right = table.rows[0].cells
    if model.get("photo") and PHOTO.exists():
        left.width = Mm(32)
        left.paragraphs[0].add_run().add_picture(str(PHOTO), width=Mm(27))
    else:
        left.width = Mm(2)
    right.width = Mm(146)
    name = right.paragraphs[0].add_run(header["name"])
    name.bold = True
    name.font.size = Pt(20)
    name.font.color.rgb = accent
    contact = right.add_paragraph(f"{header['location']} · {header['phone']} · {header['email']}\n{header['line']}")
    contact.runs[0].font.size = Pt(10)

    def heading(text_value: str):
        paragraph = document.add_paragraph()
        run = paragraph.add_run(text_value.upper())
        run.bold = True
        run.font.size = Pt(10.5)
        run.font.color.rgb = accent
        paragraph.paragraph_format.space_before = Pt(9)
        paragraph.paragraph_format.space_after = Pt(2)
        border = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for key, value in (("w:val", "single"), ("w:sz", "4"), ("w:space", "1"), ("w:color", "CCCCCC")):
            bottom.set(qn(key), value)
        border.append(bottom)
        paragraph._p.get_or_add_pPr().append(border)

    def entry_line(title: str, dates: str, meta: str):
        paragraph = document.add_paragraph()
        paragraph.paragraph_format.space_after = Pt(0)
        run = paragraph.add_run(title)
        run.bold = True
        if dates:
            paragraph.add_run("\t" + dates).font.color.rgb = RGBColor(0x55, 0x55, 0x55)
            paragraph.paragraph_format.tab_stops.add_tab_stop(Mm(180), alignment=2)
        if meta:
            meta_paragraph = document.add_paragraph(meta)
            meta_paragraph.paragraph_format.space_after = Pt(1)
            meta_paragraph.runs[0].font.color.rgb = RGBColor(0x44, 0x44, 0x44)

    for block in model["sections"]:
        heading(block["title"])
        if block["type"] == "summary":
            document.add_paragraph(block["text"])
        elif block["type"] == "experience":
            for entry in block["entries"]:
                note = f" ({entry['note']})" if entry.get("note") else ""
                entry_line(entry["title"], entry["dates"], f"{entry['employer']}{note} · {entry['place']}")
                for bullet in entry["bullets"]:
                    item = document.add_paragraph(bullet["text"], style="List Bullet")
                    item.paragraph_format.space_after = Pt(0)
        elif block["type"] == "education":
            for entry in block["entries"]:
                entry_line(entry["credential"], entry["dates"], entry["school"] + (f" · {entry['place']}" if entry.get("place") else ""))
                for line in entry.get("lines", []):
                    item = document.add_paragraph(line, style="List Bullet")
                    item.paragraph_format.space_after = Pt(0)
        else:
            document.add_paragraph(" · ".join(i["label"] for i in block["items"]))
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def cover_docx_bytes(model: Dict, job: Dict) -> bytes:
    from docx import Document
    from docx.shared import Pt
    document = Document()
    document.styles["Normal"].font.name = "Georgia"
    document.styles["Normal"].font.size = Pt(11)
    document.add_paragraph(f"{model['name']}\n{model['contact']}")
    document.add_paragraph(f"Re: {job.get('title') or ''}, {job.get('company') or ''}")
    document.add_paragraph(model["greeting"])
    for paragraph in model["paragraphs"]:
        document.add_paragraph(paragraph["text"])
    document.add_paragraph(f"{model['signoff']}\n{model['name']}")
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def chrome_path() -> Optional[str]:
    for candidate in (os.environ.get("CHROME_PATH"), "/Applications/Google Chrome.app/Contents/MacOS/Google Chrome", shutil.which("google-chrome"), shutil.which("chromium"), shutil.which("chromium-browser")):
        if candidate and Path(candidate).exists():
            return candidate
    return None


def pdf_bytes(html_value: str) -> Optional[bytes]:
    chrome = chrome_path()
    if not chrome:
        return None
    import time
    with tempfile.TemporaryDirectory() as folder:
        page = Path(folder) / "doc.html"
        target = Path(folder) / "doc.pdf"
        page.write_text(html_value)
        process = subprocess.Popen(
            [chrome, "--headless", "--disable-gpu", "--no-sandbox", "--no-first-run", "--no-default-browser-check", "--disable-extensions",
             "--no-pdf-header-footer", "--print-to-pdf-no-header", f"--user-data-dir={folder}/profile", f"--print-to-pdf={target}", page.as_uri()],
            stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL,
        )
        deadline = time.time() + 45
        last_size = -1
        while time.time() < deadline:
            if process.poll() is not None and target.exists():
                break
            if target.exists():
                size = target.stat().st_size
                if size > 0 and size == last_size:
                    break
                last_size = size
            time.sleep(0.5)
        if process.poll() is None:
            process.kill()
            process.wait(timeout=5)
        return target.read_bytes() if target.exists() and target.stat().st_size > 0 else None


def checksum(blob: bytes) -> str:
    return hashlib.sha256(blob).hexdigest()[:16]
